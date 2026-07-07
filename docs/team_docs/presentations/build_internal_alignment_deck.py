#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成「成员培训：因子 / 模型 / 数据对齐」16:9 PPT（优化版：讲者备注 + 叙事线 + 流程与收束）

用法:  python build_internal_alignment_deck.py
输出:  QuantSociety_内部培训_因子与模型对齐.pptx

联署条：将图片保存为 logos/slide_header_cobrand.png 后重跑，否则使用 logos/_generated_cobrand_strip.png
"""
from __future__ import annotations

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "QuantSociety_内部培训_因子与模型对齐.pptx")
LOGO_PRIMARY = os.path.join(HERE, "logos", "slide_header_cobrand.png")
LOGO_FALLBACK = os.path.join(HERE, "logos", "_generated_cobrand_strip.png")

BLUE = RGBColor(0, 51, 102)
BLUE_LIGHT = RGBColor(0, 76, 130)
GRAY_TXT = RGBColor(30, 36, 48)
MUTED = RGBColor(90, 98, 110)
ZEBRA = RGBColor(245, 247, 250)
GRAY_BG = RGBColor(248, 249, 251)
GRAY_LINE = RGBColor(200, 204, 210)
CALLOUT = RGBColor(0, 92, 140)


def ensure_logo_png() -> str:
    if os.path.isfile(LOGO_PRIMARY):
        return LOGO_PRIMARY
    os.makedirs(os.path.join(HERE, "logos"), exist_ok=True)
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("需要 Pillow: pip install Pillow", file=sys.stderr)
        return ""
    w, h = 1400, 110
    im = Image.new("RGB", (w, h), (255, 255, 255))
    dr = ImageDraw.Draw(im)
    for x in (w // 3, 2 * w // 3):
        dr.line([(x, 12), (x, h - 12)], fill=(180, 186, 194), width=1)
    try:
        f_title = ImageFont.truetype("DejaVuSans.ttf", 14)
        f_small = ImageFont.truetype("DejaVuSans.ttf", 12)
    except OSError:
        f_title = f_small = ImageFont.load_default()
    dr.text((20, 12), "THE HONG KONG\nUNIVERSITY OF SCIENCE\nAND TECHNOLOGY", fill=(0, 51, 102), font=f_title)
    dr.text((w // 3 + 20, 28), "MSc IN\nFINANCIAL MATHEMATICS", fill=(0, 51, 102), font=f_title)
    cx2 = 2 * w // 3 + 24
    dr.text((cx2, 20), "QUANT", fill=(0, 51, 102), font=f_title)
    dr.text((cx2 + 2, 58), "s o c i e t y", fill=(0, 76, 130), font=f_small)
    im.save(LOGO_FALLBACK, "PNG")
    return LOGO_FALLBACK


def set_widescreen(prs: Presentation) -> None:
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)


def add_top_accent_bar(slide) -> None:
    r = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.06)
    )
    r.fill.solid()
    r.fill.fore_color.rgb = BLUE
    r.line.fill.background()


def add_top_right_logo(slide, prs: Presentation, path: str) -> None:
    if not path or not os.path.isfile(path):
        return
    w, h = Inches(4.35), Inches(0.65)
    left = prs.slide_width - w - Inches(0.18)
    top = Inches(0.14)
    slide.shapes.add_picture(path, left, top, width=w, height=h)


def set_notes(slide, text: str) -> None:
    ns = slide.notes_slide
    ns.notes_text_frame.text = text.strip()


def _style_title(tf, size=Pt(24)) -> None:
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.bold = True
        r.font.size = size
        r.font.name = "Microsoft YaHei"
        r.font.color.rgb = BLUE


def _style_bullets(tf, size=Pt(15), bold_lead: bool = False) -> None:
    for p in tf.paragraphs:
        p.space_after = Pt(8)
        p.alignment = PP_ALIGN.LEFT
        t = p.text
        if hasattr(p, "clear"):
            p.clear()  # python-pptx 较新版本
        else:
            p.text = ""
        if bold_lead and "：" in t:
            a, b = t.split("：", 1)
            r0 = p.add_run()
            r0.text = a + "："
            r0.font.bold = True
            r0.font.size = size
            r0.font.name = "Microsoft YaHei"
            r0.font.color.rgb = BLUE
            r1 = p.add_run()
            r1.text = b
        else:
            r1 = p.add_run()
            r1.text = t
        r1.font.size = size
        r1.font.name = "Microsoft YaHei"
        r1.font.color.rgb = GRAY_TXT


def add_footer_bar(slide, text: str) -> None:
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.75), Inches(13.333), Inches(0.75)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = GRAY_BG
    sh.line.fill.background()
    if not text:
        return
    tf = sh.text_frame
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        p.margin_left = Inches(0.35)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.name = "Microsoft YaHei"
            r.font.color.rgb = MUTED


def add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    prs_w,
    logo_path: str,
    foot: str = "",
    notes: str = "",
    callout: str = "",
    y_body: float = 0.95,
    bold_lead: bool = True,
) -> object:
    li = 6 if len(prs.slide_layouts) > 6 else 5
    blank = prs.slides.add_slide(prs.slide_layouts[li])
    add_top_accent_bar(blank)
    add_footer_bar(blank, foot)

    tbox = blank.shapes.add_textbox(Inches(0.45), Inches(0.4), Inches(8.2), Inches(0.7))
    tbox.text_frame.text = title
    _style_title(tbox.text_frame, Pt(24))

    if callout:
        cb = blank.shapes.add_textbox(Inches(0.45), Inches(0.82), Inches(12.0), Inches(0.38))
        cb.text_frame.text = callout
        cb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        cb.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
        cb.text_frame.paragraphs[0].runs[0].font.color.rgb = CALLOUT
        cb.text_frame.paragraphs[0].runs[0].font.italic = True
        y_body = 1.2

    box = blank.shapes.add_textbox(Inches(0.45), Inches(y_body), Inches(12.0), Inches(5.35))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = bullets[0] if bullets else ""
    for line in bullets[1:]:
        p = tf.add_paragraph()
        p.text = line
    _style_bullets(tf, Pt(15), bold_lead=bold_lead)
    add_top_right_logo(blank, prs_w, logo_path)
    if notes:
        set_notes(blank, notes)
    return blank


def add_two_column_slide(
    prs,
    title: str,
    left_title: str,
    left_bullets: list[str],
    right_title: str,
    right_bullets: list[str],
    prs_w,
    logo_path: str,
    foot: str = "",
    notes: str = "",
) -> object:
    li = 6 if len(prs.slide_layouts) > 6 else 5
    s = prs.slides.add_slide(prs.slide_layouts[li])
    add_top_accent_bar(s)
    add_footer_bar(s, foot)
    tbox = s.shapes.add_textbox(Inches(0.45), Inches(0.38), Inches(10.0), Inches(0.65))
    tbox.text_frame.text = title
    _style_title(tbox.text_frame, Pt(22))

    wcol = 5.9
    # 左
    lt = s.shapes.add_textbox(Inches(0.4), Inches(0.95), Inches(wcol), Inches(0.35))
    lt.text_frame.text = left_title
    lt.text_frame.paragraphs[0].runs[0].font.bold = True
    lt.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    lt.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    lt.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
    lb = s.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(wcol), Inches(4.7))
    ltf = lb.text_frame
    ltf.text = left_bullets[0] if left_bullets else ""
    for x in left_bullets[1:]:
        ltf.add_paragraph().text = x
    _style_bullets(ltf, Pt(13.5), bold_lead=True)
    # 右
    rt = s.shapes.add_textbox(Inches(6.65), Inches(0.95), Inches(wcol), Inches(0.35))
    rt.text_frame.text = right_title
    rt.text_frame.paragraphs[0].runs[0].font.bold = True
    rt.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    rt.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    rt.text_frame.paragraphs[0].runs[0].font.color.rgb = CALLOUT
    rb = s.shapes.add_textbox(Inches(6.65), Inches(1.3), Inches(wcol), Inches(4.7))
    rtf = rb.text_frame
    rtf.text = right_bullets[0] if right_bullets else ""
    for x in right_bullets[1:]:
        rtf.add_paragraph().text = x
    _style_bullets(rtf, Pt(13.5), bold_lead=True)
    # 中间竖线
    ln = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.45), Inches(0.9), Inches(0.02), Inches(5.2)
    )
    ln.fill.solid()
    ln.fill.fore_color.rgb = GRAY_LINE
    ln.line.fill.background()
    add_top_right_logo(s, prs_w, logo_path)
    if notes:
        set_notes(s, notes)
    return s


def add_flow_pipeline_slide(prs, logo_path: str, notes: str) -> object:
    li = 6 if len(prs.slide_layouts) > 6 else 5
    s = prs.slides.add_slide(prs.slide_layouts[li])
    add_top_accent_bar(s)
    tbox = s.shapes.add_textbox(Inches(0.45), Inches(0.38), Inches(10.0), Inches(0.65))
    tbox.text_frame.text = "一条链：从「多引擎」到「下游能消费」"
    _style_title(tbox.text_frame, Pt(22))
    st = s.shapes.add_textbox(Inches(0.45), Inches(0.9), Inches(12.2), Inches(0.45))
    st.text_frame.text = "（与 04 篇一 资产流 同构；会上一分钟走一遍，细节会后看文档。）"
    st.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    st.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    st.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
    # 五框 + 四箭头
    labels = [
        "多引擎\n(A/B/C/D)",
        "候选池 P",
        "五安检 G",
        "注册 R\n元数据",
        "下游 D\n03/02",
    ]
    n = len(labels)
    box_w, box_h, gap, y0 = Inches(1.95), Inches(0.85), Inches(0.28), Inches(2.25)
    total_w = n * float(box_w) + (n - 1) * (float(gap) + 0.35)  # 约
    x0 = (float(Inches(13.333)) - (n * float(box_w) + (n - 1) * 0.45 * 914400 / 12700)) / 914400 * 12700
    # 简化为从左边排
    x = 0.35
    centers = []
    for i, lab in enumerate(labels):
        left = Inches(x)
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y0, box_w, box_h)
        try:
            sh.adjustments[0] = 0.12
        except (AttributeError, TypeError, ValueError):
            pass
        if i in (2, 3):
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(230, 240, 250)
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sh.line.color.rgb = BLUE_LIGHT
        sh.line.width = Pt(1.25)
        tf = sh.text_frame
        tf.text = lab
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        for r in tf.paragraphs[0].runs:
            r.font.name = "Microsoft YaHei"
            r.font.size = Pt(12)
            r.font.color.rgb = BLUE
        centers.append(x + 1.95 / 2)
        x += 1.95 + 0.45
    # 箭头 简化为 "→" 形状之间文字
    for i in range(n - 1):
        ax = 0.35 + (i + 1) * (1.95 + 0.45) - 0.4
        arr = s.shapes.add_textbox(Inches(ax), Inches(2.45), Inches(0.4), Inches(0.4))
        arr.text_frame.text = "→"
        arr.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        arr.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
        arr.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
    # 底注
    add_footer_bar(
        s,
        "口播落点：四引擎 只增加发现率，不自动提高入库质量；G 是统一闸门。",
    )
    add_top_right_logo(s, prs, logo_path)
    set_notes(s, notes)
    return s


def add_table_slide_zebra(
    prs, title: str, headers: list[str], rows: list[list[str]], prs_w, logo_path: str, notes: str
):
    li = 6 if len(prs.slide_layouts) > 6 else 5
    blank = prs.slides.add_slide(prs.slide_layouts[li])
    add_top_accent_bar(blank)
    tbox = blank.shapes.add_textbox(Inches(0.4), Inches(0.32), Inches(11.0), Inches(0.55))
    tbox.text_frame.text = title
    _style_title(tbox.text_frame, Pt(20))
    rrows, cols = len(rows) + 1, len(headers)
    tbl = blank.shapes.add_table(
        rrows, cols, Inches(0.32), Inches(0.95), Inches(12.65), Inches(4.55)
    ).table
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = BLUE
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.name = "Microsoft YaHei"
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            c = tbl.cell(i, j)
            c.text = cell
            if i % 2 == 0:
                c.fill.solid()
                c.fill.fore_color.rgb = ZEBRA
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.name = "Microsoft YaHei"
    add_footer_bar(blank, "讲法：只要求同事记住本表「第一列+第三列各一个词」。第二列会後自读 04 篇二 2.0.2。")
    add_top_right_logo(blank, prs_w, logo_path)
    if notes:
        set_notes(blank, notes)
    return blank


def add_compact_matrix_slide(prs, logo_path, notes) -> object:
    """5×4 简写：只显示「关键红线」 缩写，避免整墙字。"""
    li = 6 if len(prs.slide_layouts) > 6 else 5
    s = prs.slides.add_slide(prs.slide_layouts[li])
    add_top_accent_bar(s)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(11.5), Inches(0.6))
    tb.text_frame.text = "五安检 × 四引擎：一眼看「哪格最容易着」"
    _style_title(tb.text_frame, Pt(20))
    sub = s.shapes.add_textbox(Inches(0.4), Inches(0.82), Inches(12.0), Inches(0.4))
    sub.text_frame.text = "格内是「最先要问的那一句」 不是完整定义；全表见 04 执行摘要 矩阵。"
    sub.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    sub.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    sub.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
    headers = ["", "A 演化", "B RL", "C LLM", "D 信息/拓扑"]
    data = [
        ["安检1 假发现", "膨胀/试次", "隐式多试", "默写/幻觉", "多窗 p-hack"],
        ["安检2 摩擦", "高换手式", "Δ 奖→频", "未估换手", "算力/评估"],
        ["安检3 正交", "同质化", "模式崩溃", "同构", "被特征绑架"],
        ["安检4 体制", "单窗", "非稳", "外推", "薄窗"],
        ["安检5 白盒", "量纲", "黑箱", "对不上", "难翻译"],
    ]
    r, c = len(data) + 1, len(headers)
    table = s.shapes.add_table(r, c, Inches(0.32), Inches(1.2), Inches(12.65), Inches(4.2)).table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if j else RGBColor(60, 80, 100)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for ru in p.runs:
                ru.font.size = Pt(9)
                ru.font.name = "Microsoft YaHei"
                ru.font.bold = True
                ru.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(data, 1):
        for j, t in enumerate(row):
            cell = table.cell(i, j)
            cell.text = t
            if i % 2 == 0 and j:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ZEBRA
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
                for ru in p.runs:
                    ru.font.size = Pt(8.5)
                    ru.font.name = "Microsoft YaHei"
    add_footer_bar(
        s,
        "口播 30 秒：我们不是在找「能写多炫」，是在找「哪格红线最先亮」。",
    )
    add_top_right_logo(s, prs, logo_path)
    set_notes(s, notes)
    return s


def build():
    logo = ensure_logo_png()
    prs = Presentation()
    set_widescreen(prs)

    # ---- 0 标题
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = "因子 × 模型 × 数据：内部对齐"
    s0.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    s0.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    s0.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
    if len(s0.shapes) > 1 and s0.shapes[1].has_text_frame:
        s0.shapes[1].text = (
            "QUANT society · 成员培训\n"
            "本 deck：讲清边界、闸门与对接 03 的一句话 · 不替代仓库长文 01/02/03/04"
        )
        s0.shapes[1].text_frame.paragraphs[0].runs[0].font.size = Pt(15)
    add_top_right_logo(s0, prs, logo)
    set_notes(
        s0,
        "【开场 60 秒】先报目的：不是 Code walkthrough，是让大家会后能「把人找对、把表填对、把事问到点子上」。\n"
        "停 1 秒：听的人里谁主因子/谁主模型/谁主数据，心里各选一个动作（最后 Q&A 用）。\n"
        "过渡：先给今天三张片子就能跟完全程——议程。",
    )

    # ---- 1 议程
    add_content_slide(
        prs,
        "今天 25 分钟怎么听（可压缩到 15）",
        [
            "1）边界：数据 / 因子 / 模型 / 策略 各停在哪一段「可辩护」",
            "2）闸门：四引擎只是搜索偏置，多引擎合流进同一套五安检，再进注册与下游",
            "3）对接 03：先 A/B/C 与「一行是什么」，再谈损失与 WFO/八续；上游泄漏没有「后处理」",
            "（可选加钟）看五安检×四引擎矩阵：认格子，不背定义",
        ],
        prs,
        logo,
        foot="时间不够：跳过「矩阵」与「A/B/C 小表」两张，收束三句 仍要讲。",
        notes="【讲法】手比「三指」对应三块。说清：细节在文档，会上一致口径。\n"
        "若听众一半工程一半研究：加一句「论文细节会后」 即可压住跑题。",
        callout="给听众的抓手：你只需决定——会后打开 04 还是 03 的哪一篇。",
    )

    # ---- 2 金线
    add_content_slide(
        prs,
        "本组在比什么（一句，写在心头）",
        [
            "比的不是「出多少条因子」，是「在统一闸门下，少漏 进脏信号」",
            "多引擎＝在搜索层叠不同偏置；「不」 等于少一道 统计与工程纪律",
            "模型组 03 的 OOS 故事，在因子带系统性标签/时间错误时 无 可靠前提——这点必须当场听懂",
        ],
        prs,
        logo,
        notes="【重音】第二句常有人忘：引擎多≠安检少。\n"
        "第三句：有新人眼神飘——举一个「用未来行业」 例子停 2 秒，不展开公式。",
        callout="讲者自检：全 room 能复述「多引擎不减税」 即过关。",
    )

    # ---- 3 该放/不放（两栏）
    add_two_column_slide(
        prs,
        "会上讲 / 会下自己啃",
        "会上要带走",
        [
            "可画出来的责任段：谁为 PiT/可得、谁为因子候选、谁为 WFO&契约、谁为成本",
            "可复述的链：多引擎 → P → G → R → D（能背五个字母顺序即可）",
            "五安检名字 + 每类「一句验收」：会后对着 04/篇十四 打勾",
        ],
        "会上故意不讲（避免念稿）",
        [
            "具体路径/类名/YAML/命令行（放 01/02/03/04 + 合入说明）",
            "未内审的公开展示夏普/IC 数字作论据",
            "WFO/八续/MLflow 的数学细节（在 03 篇八+八续）",
        ],
        prs,
        logo,
        foot="新人若追问命令：答「在 篇 1.5 SOP 与合入说明」；不在这 room 里打开终端。",
        notes="【互动】可问半句：「我们上周最容易死在 G 的哪一条？」 不冷场就接 五安检 页。",
    )

    # ---- 4 流程
    add_flow_pipeline_slide(
        prs,
        logo,
        notes="【手比】从左指到右。停到 G：强调「不替代」——引擎只是入口。\n"
        "停到 R：说清「可注册=元数据齐」。停到 D：点 03+02，不展开。",
    )

    # ---- 5 责任 四段
    add_content_slide(
        prs,
        "各段责任（与 01/02/03/04 对表，不讲路径）",
        [
            "01：登记、发布、读数 一致；PiT/可得日 的 事实源",
            "04 因子组：候选+五安检+注册 材料；不背「帮模型把 OOS 拉好看」 的锅",
            "03 模型组：WFO+试次+信号契约 下训练/交付；吃池方式 写在 03",
            "02/策略：可交易、成本、从信号到仓位 的 主叙事 不在 03/04 包圆",
        ],
        prs,
        logo,
        notes="【常见反驳】若有人说「我们因子好了模型自然好」：用第三行「契约与 WFO 仍在 03」 回应。\n"
        "不展开「谁写 PR」——踢回 02。",
    )

    # ---- 6 四引擎
    add_table_slide_zebra(
        prs,
        "四引擎：你只需记住「偏置 + 最危险线」",
        ["引擎", "听会抓这句", "最先爆的红线（点到为止）"],
        [
            [
                "A 符号/演化",
                "文法里搜式子、复杂度罚+登记",
                "膨胀/同质/唯 in-sample 好看",
            ],
            [
                "B 协同/序列",
                "奖励看组合边际、防单 IC 自嗨",
                "非稳/模式崩/隐式多试次",
            ],
            [
                "C LLM/Agent",
                "话必须落到可执行式+可复现回测",
                "默写/叙事幻觉/结构距过近",
            ],
            [
                "D 信息/拓扑",
                "估计贵、先 研究档+成本",
                "多窗/重算=隐性多重比较",
            ],
        ],
        prs,
        logo,
        notes="【节奏】每行 15 秒。停 B：举例「动量都高 IC」 仍可边际差——衔接「要组合奖励」。\n"
        "停 C：强调「能跑」 不是「能入库」。",
    )

    # ---- 7 五安检
    add_content_slide(
        prs,
        "五安检：会后能对着清单打勾",
        [
            "1 假发现/复杂度：禁「唯 in-sample」；登记、锁箱、试次/窗长 有上限",
            "2 摩擦：看 净；高换手纸面 不算交付",
            "3 正交/增量：对 核心库 的 边际/残差；禁红海 凑数",
            "4 体制/尾部：压测+半衰/拥挤 至少能写 一句 处置或标签",
            "5 白盒/量纲：经济意义+量纲/语法 在登记里 可指",
        ],
        prs,
        logo,
        foot="主持人可在此停 10 秒：问「我们当前最缺哪一条材料？」 记到会议纪要。",
        notes="【不要逐条念定义】用「能验收」 三个字作节拍。\n"
        "若时间紧：只展开 1 与 3，其余说「见 04 篇四～五」。",
    )

    # ---- 8 矩阵
    add_compact_matrix_slide(
        prs,
        logo,
        notes="【可跳过】时间不够整页 30 秒：只指左上角三格：假发现/摩擦/正交。\n"
        "用于研究同学「找自己的坑」 而不是管理汇报。",
    )

    # ---- 9 任务 A B C
    add_table_slide_zebra(
        prs,
        "先定「一行数据是什么」（03 篇三；会上只记三格）",
        ["类型", "行/张量 怎么说", "和今天的关系"],
        [
            [
                "A 截面/日",
                "日×票 矩阵；多排序/分组",
                "因子多引擎 多数先落 A 的评估口径",
            ],
            [
                "B 时序",
                "票×时间步 序列；沿时间 为主",
                "损失与 WFO 不能按 A 的口播 偷换",
            ],
            [
                "C 混合",
                "两阶写清：各自损失/窗/验证",
                "禁止口头「混合」 而无结构图 或表",
            ],
        ],
        prs,
        logo,
        notes="【桥接 03】念完立刻接下一页：「定类了再谈 篇八/八续」。\n"
        "有人提问「我们算哪种」：答「以试验说明第一行 为准，争议找负责人」。",
    )

    # ---- 10 对接 03
    add_content_slide(
        prs,
        "对接模型组：三句，展开都在 03",
        [
            "一句顺序：先 定 A/B/C 与「行定义」→再 定损失 与 篇八 WFO/八续 主叙事",
            "一句硬边界：因子 信息集/标签/宇宙 的 系统性 错误 = 下游 没有 可辩护 OOS 前提，不是多训能救",
            "一句交付物：过闸 的 可注册 因子 + 元数据；不是裸 series",
        ],
        prs,
        logo,
        notes="【收束前高潮】三句 各读一遍 慢 半拍。\n"
        "可提问：「我们上周哪个因子 缺 PiT 字段？」 把会开到地上。",
        callout="与工程对齐：开 issue 时贴「A/B/C + 行定义」 在描述首行。",
    )

    # ---- 11 自读
    add_content_slide(
        prs,
        "会后自读（只选一条最深就行）",
        [
            "走因子/挖掘：04 卷 A → 篇二 2.0.1/2.0.2 → 篇四～五 → 篇十四 勾 1/3 项",
            "走训练/主仓：03 篇 1.5 与 1.5.2a/2b、篇三 3.2.5/5a、篇八+八续",
            "走数据：01 与 因子 登记/PiT 对表 字段；和 04 篇七 时间律 对齐",
        ],
        prs,
        logo,
        foot="发材料：Slack/邮件 附本 deck + 两白皮书文件名；不要附整仓 zip。",
        notes="【防 overwhelm】说「三选一，不要全读」—— 降低沉默成本。",
    )

    # ---- 12 收束
    add_content_slide(
        prs,
        "收束三句（可当散场口播）",
        [
            "多引擎 解决「找得到」，五安检+注册 解决「敢用」；两回事",
            "和 03 说话的前缀永远是：A/B/C + 行定义；其它都是后话",
            "今天的作业：各角色各打开一篇长文 约 10 分钟，把「最卡你的一行」 记在会议纪要",
        ],
        prs,
        logo,
        notes="【结尾】不 Q&A 也可三句 结束。需要气势：最后一句 提高半度音量「写进纪要」。\n"
        "若 CEO/顾问在场：不加班数据，不承诺未跑通的数字。",
    )

    # ---- 13 Q&A
    s_end = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else 5)
    add_top_accent_bar(s_end)
    t = s_end.shapes.add_textbox(Inches(0.5), Inches(2.35), Inches(8), Inches(1.0))
    t.text_frame.text = "Q & A"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    t.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
    t.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
    st = s_end.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(11.5), Inches(1.5))
    st.text_frame.text = (
        "会后可约小桌：\n"
        "· 试验登记 / 锁箱 / 试次 预算 对表\n"
        "· OOS 主指标 与 组合层 接口 对表\n"
        "· 与本次 deck 对不上的旧流程，走「一次 纠偏会」 不要群里碎片化吵"
    )
    for p in st.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(16)
            r.font.name = "Microsoft YaHei"
    add_footer_bar(s_end, "感谢 · 本文件可由 build_internal_alignment_deck.py 重新生成以同步叙事修订。")
    add_top_right_logo(s_end, prs, logo)
    set_notes(
        s_end,
        "【Q&A 引导】没问题时用「那我们从 PiT 字段开始 每人说一条」 破冰。\n"
        "收束问题类型：去 01/ 去 04/ 去 03/ 要负责人决策——四类分流。",
    )

    prs.save(OUT)
    print("Wrote", OUT, "slides:", len(prs.slides))
    if logo:
        print("Logo:", logo)


if __name__ == "__main__":
    build()
